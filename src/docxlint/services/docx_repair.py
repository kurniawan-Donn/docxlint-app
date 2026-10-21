"""Selective DOCX repair engine Phase 6.

Keputusan FINAL yang dipatuhi:
- DECISION-002: rPr disalin dengan copy.deepcopy.
- DECISION-004: apply order offset menurun.
- DECISION-005: overlap -> conflict, tidak otomatis.
- DECISION-006: stale check per finding.
- DECISION-013/023: rollback berbasis byte DOCX asli.
- DECISION-014: unduh diblokir jika ada failed kritis.
- DECISION-021: run kompleks -> skipped.
- DECISION-024: failed kritis = body/header/footer/tabel.

Original DOCX tidak pernah di-overwrite. Edit konservatif:
- replace_text: hanya dalam satu run sederhana.
- apply_italic: hanya dalam satu run sederhana; split run diizinkan
  dengan deepcopy rPr.
- Cross-run -> unsupported.
"""

import copy
import hashlib
import uuid
from io import BytesIO
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..models.change_plan import (
    APPLIED,
    CONFLICT,
    ChangeItem,
    ChangePlan,
    FAILED,
    NOT_APPLIED,
    OPERATION_APPLY_ITALIC,
    OPERATION_REPLACE_TEXT,
    REJECTED,
    SKIPPED,
    STALE,
    UNSUPPORTED,
    VALID,
)
from ..models.document import DocumentSnapshot
from ..models.finding import Finding


def compute_snapshot_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# --- resolusi block_path -> paragraf python-docx -------------------------

def _resolve_block(doc, block_path: str):
    parts = block_path.split("/")
    root = parts[0]
    if root == "body":
        node = doc
    elif root.startswith("header[") or root.startswith("footer["):
        idx = int(root[root.index("[") + 1:root.index("]")])
        section = doc.sections[idx]
        node = section.header if root.startswith("header[") else section.footer
    else:
        raise ValueError(f"unknown root: {root}")

    i = 1
    while i < len(parts):
        token = parts[i]
        if token.startswith("p["):
            node = node.paragraphs[int(token[2:-1])]
        elif token.startswith("table["):
            node = node.tables[int(token[6:-1])]
        elif token.startswith("row["):
            node = node.rows[int(token[4:-1])]
        elif token.startswith("cell["):
            node = node.cells[int(token[5:-1])]
        else:
            raise ValueError(f"unknown segment: {token}")
        i += 1
    return node


# --- deteksi run sederhana -----------------------------------------------

def _is_simple_run_element(run_el) -> bool:
    t_count = 0
    for child in run_el:
        tag = child.tag
        if tag == qn("w:rPr"):
            continue
        if tag == qn("w:t"):
            t_count += 1
            if t_count > 1:
                return False
            continue
        return False
    return True


class _Skip(Exception):
    def __init__(self, reason: str, status: str = SKIPPED) -> None:
        super().__init__(reason)
        self.reason = reason
        self.status = status


def _replace_in_run(paragraph, run_index: int,
                    char_start: int, char_end: int,
                    original_value: str, replacement: str) -> None:
    runs = paragraph.runs
    if run_index < 0 or run_index >= len(runs):
        raise _Skip("run_index_out_of_range", FAILED)
    run = runs[run_index]
    run_el = run._element
    if not _is_simple_run_element(run_el):
        raise _Skip("run_complex", SKIPPED)

    text = run.text or ""
    if char_start < 0 or char_end > len(text) or char_start >= char_end:
        raise _Skip("offset_out_of_range", FAILED)
    actual = text[char_start:char_end]
    if actual != original_value:
        raise _Skip("original_value_mismatch", STALE)

    new_text = text[:char_start] + replacement + text[char_end:]

    rpr = run_el.find(qn("w:rPr"))
    for t in run_el.findall(qn("w:t")):
        run_el.remove(t)
    new_t = OxmlElement("w:t")
    new_t.text = new_text
    new_t.set(qn("xml:space"), "preserve")
    if rpr is not None:
        rpr.addnext(new_t)
    else:
        run_el.insert(0, new_t)


def _apply_italic_in_run(paragraph, run_index: int,
                         char_start: int, char_end: int,
                         original_value: str) -> None:
    """Apply italic ke range dalam satu run.

    Jika range = seluruh run: set italic pada rPr run.
    Jika range < seluruh run: split menjadi 3 run dengan rPr deepcopy.
    """
    runs = paragraph.runs
    if run_index < 0 or run_index >= len(runs):
        raise _Skip("run_index_out_of_range", FAILED)
    run = runs[run_index]
    run_el = run._element
    if not _is_simple_run_element(run_el):
        raise _Skip("run_complex", SKIPPED)

    text = run.text or ""
    if char_start < 0 or char_end > len(text) or char_start >= char_end:
        raise _Skip("offset_out_of_range", FAILED)
    actual = text[char_start:char_end]
    if actual != original_value:
        raise _Skip("original_value_mismatch", STALE)

    parent = run_el.getparent()
    rpr = run_el.find(qn("w:rPr"))
    rpr_copy = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
    # Pastikan italic=True di rpr_copy
    for tag in ("w:i", "w:iCs"):
        for el in rpr_copy.findall(qn(tag)):
            rpr_copy.remove(el)
    i_el = OxmlElement("w:i")
    rpr_copy.append(i_el)

    def _make_run(segment_text: str, italic: bool):
        r = OxmlElement("w:r")
        r_rpr = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
        if italic:
            r_rpr = copy.deepcopy(rpr_copy)
        if len(r_rpr) > 0:
            r.append(r_rpr)
        t = OxmlElement("w:t")
        t.text = segment_text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        return r

    before = text[:char_start]
    target = text[char_start:char_end]
    after = text[char_end:]

    new_elements = []
    if before:
        new_elements.append(_make_run(before, False))
    new_elements.append(_make_run(target, True))
    if after:
        new_elements.append(_make_run(after, False))

    idx = list(parent).index(run_el)
    parent.remove(run_el)
    for offset, el in enumerate(new_elements):
        parent.insert(idx + offset, el)


# --- engine ---------------------------------------------------------------

class RepairEngine:
    def __init__(self, original_bytes: bytes,
                 snapshot: DocumentSnapshot) -> None:
        self.original_bytes = original_bytes
        self.snapshot = snapshot
        self.snapshot_hash = compute_snapshot_hash(original_bytes)

    # --- build plan ----------------------------------------------------

    def build_change_plan(
        self,
        findings: List[Finding],
        approved_actions: Dict[str, dict],
    ) -> ChangePlan:
        """approved_actions: {finding_id: {"operation": ..., "replacement": ...}}"""
        plan = ChangePlan(
            snapshot_id=self.snapshot.snapshot_id,
            snapshot_hash=self.snapshot_hash,
        )
        for f in findings:
            action = approved_actions.get(f.finding_id)
            if not action:
                continue
            op = action.get("operation")
            replacement = action.get("replacement")
            item = ChangeItem(
                change_id=uuid.uuid4().hex,
                finding_id=f.finding_id,
                operation=op,
                part=f.location.part,
                block_path=f.location.block_path,
                run_index=f.location.run_index,
                char_start=f.location.char_start,
                char_end=f.location.char_end,
                original_value=f.original_value,
                replacement=replacement,
                snapshot_hash=self.snapshot_hash,
                global_text_offset=f.location.global_text_offset,
            )
            plan.add(item)
        return plan

    # --- validate ------------------------------------------------------

    def validate_plan(self, plan: ChangePlan) -> ChangePlan:
        # Snapshot hash cocok?
        if plan.snapshot_hash != self.snapshot_hash:
            for item in plan.items:
                item.validation = STALE
                item.reason = "snapshot_hash_mismatch"
            return plan

        # Unsupported operations?
        for item in plan.items:
            if item.operation not in (
                OPERATION_REPLACE_TEXT, OPERATION_APPLY_ITALIC,
            ):
                item.validation = UNSUPPORTED
                item.reason = "unknown_operation"
            elif item.operation == OPERATION_REPLACE_TEXT and \
                    item.replacement is None:
                item.validation = REJECTED
                item.reason = "missing_replacement"

        # Deteksi overlap per (part, block_path, run_index).
        groups: Dict[Tuple[str, str, int], List[ChangeItem]] = {}
        for item in plan.items:
            if item.validation != VALID:
                continue
            key = (item.part, item.block_path, item.run_index)
            groups.setdefault(key, []).append(item)

        for key, items in groups.items():
            for i in range(len(items)):
                for j in range(i + 1, len(items)):
                    a, b = items[i], items[j]
                    if a.char_start < b.char_end and \
                            b.char_start < a.char_end:
                        a.validation = CONFLICT
                        b.validation = CONFLICT
                        a.reason = "overlap"
                        b.reason = "overlap"

        return plan

    # --- apply ---------------------------------------------------------

    def apply_plan(self, plan: ChangePlan) -> Tuple[bytes, ChangePlan]:
        """Apply valid items. Return (output_bytes, plan)."""
        valid = [i for i in plan.items if i.validation == VALID]
        if not valid:
            return self.original_bytes, plan

        try:
            doc = Document(BytesIO(self.original_bytes))
        except Exception:
            for i in valid:
                i.execution = FAILED
                i.reason = "open_failed"
            return self.original_bytes, plan

        # Order: per (part, block_path, run_index) group apply
        # descending char_start.
        groups: Dict[Tuple[str, str, int], List[ChangeItem]] = {}
        for item in valid:
            key = (item.part, item.block_path, item.run_index)
            groups.setdefault(key, []).append(item)

        for key, items in groups.items():
            items.sort(key=lambda x: -x.char_start)
            try:
                paragraph = _resolve_block(doc, key[1])
            except Exception:
                for i in items:
                    i.execution = FAILED
                    i.reason = "block_resolve_failed"
                continue
            for item in items:
                try:
                    if item.is_text_edit():
                        _replace_in_run(
                            paragraph, item.run_index,
                            item.char_start, item.char_end,
                            item.original_value, item.replacement or "",
                        )
                    elif item.is_italic_edit():
                        _apply_italic_in_run(
                            paragraph, item.run_index,
                            item.char_start, item.char_end,
                            item.original_value,
                        )
                    item.execution = APPLIED
                except _Skip as e:
                    item.execution = e.status
                    item.reason = e.reason
                except Exception:
                    item.execution = FAILED
                    item.reason = "unexpected_error"

        try:
            out = BytesIO()
            doc.save(out)
            return out.getvalue(), plan
        except Exception:
            # Rollback: kembalikan byte asli (DECISION-013/023).
            for item in plan.items:
                if item.execution == APPLIED:
                    item.execution = FAILED
                    item.reason = "save_failed_rollback"
            return self.original_bytes, plan

    # --- failed kritis -------------------------------------------------

    @staticmethod
    def has_critical_failure(plan: ChangePlan) -> bool:
        """DECISION-024: failed di body/header/footer/tabel = kritis."""
        for item in plan.items:
            if item.execution != FAILED:
                continue
            if item.part == "document" and "/table[" not in item.block_path:
                # body utama
                return True
            if item.part.startswith("header[") or item.part.startswith("footer["):
                return True
            if "/table[" in item.block_path:
                return True
        return False


__all__ = ["RepairEngine", "compute_snapshot_hash"]