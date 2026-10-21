"""Pembaca DOCX.

Membangun DocumentSnapshot dari bytes DOCX tanpa mengubah dokumen.
Tidak meratakan dokumen menjadi satu string.

Deteksi run kompleks mengikuti DECISION-021:
- lebih dari satu <w:t>
- mengandung <w:br>, <w:tab>, <w:drawing>, <w:footnoteReference>,
  <w:endnoteReference>, <w:proofErr>, <w:fldChar>, <w:instrText>
- mengandung elemen anak selain <w:rPr> dan <w:t>
"""

import uuid
from io import BytesIO
from typing import List, Optional

from docx import Document
from docx.oxml.ns import qn

from ..models.document import Block, DocumentSnapshot, Run


def _is_complex_run(run_element) -> bool:  # type: ignore[no-untyped-def]
    t_children = run_element.findall(qn("w:t"))
    if len(t_children) > 1:
        return True
    for child in run_element:
        tag = child.tag
        if tag == qn("w:rPr") or tag == qn("w:t"):
            continue
        return True
    return False


def _extract_runs(paragraph) -> List[Run]:  # type: ignore[no-untyped-def]
    runs: List[Run] = []
    offset = 0
    for idx, run in enumerate(paragraph.runs):
        text = run.text or ""
        start = offset
        end = offset + len(text)
        offset = end
        runs.append(
            Run(
                run_index=idx,
                text=text,
                char_start=start,
                char_end=end,
                is_complex_run=_is_complex_run(run._element),
                bold=run.bold,
                italic=run.italic,
                xml_element=run._element,
            )
        )
    return runs


def _read_paragraph_block(
    paragraph,  # type: ignore[no-untyped-def]
    part: str,
    block_path: str,
    block_type: str,
) -> Block:
    runs = _extract_runs(paragraph)
    text = "".join(r.text for r in runs)
    style_name = (
        paragraph.style.name if paragraph.style is not None else None
    )
    numbering = None
    pPr = paragraph._element.find(qn("w:pPr"))
    if pPr is not None:
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            ilvl = numPr.find(qn("w:ilvl"))
            numId = numPr.find(qn("w:numId"))
            numbering = {
                "ilvl": ilvl.get(qn("w:val")) if ilvl is not None else None,
                "numId": numId.get(qn("w:val")) if numId is not None else None,
            }
    return Block(
        block_path=block_path,
        block_type=block_type,
        part=part,
        text=text,
        runs=runs,
        style_name=style_name,
        numbering=numbering,
        xml_element=paragraph._element,
    )


def _classify_paragraph(paragraph) -> str:  # type: ignore[no-untyped-def]
    style_name = (
        paragraph.style.name if paragraph.style is not None else ""
    ) or ""
    if style_name.lower().startswith("heading"):
        return "heading"
    pPr = paragraph._element.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:numPr")) is not None:
        return "list_item"
    return "paragraph"


def _read_paragraphs(paragraphs, part: str, prefix: str) -> List[Block]:  # type: ignore[no-untyped-def]
    blocks: List[Block] = []
    for idx, para in enumerate(paragraphs):
        blocks.append(
            _read_paragraph_block(
                para,
                part,
                f"{prefix}/p[{idx}]",
                _classify_paragraph(para),
            )
        )
    return blocks


def _read_table(table, part: str, prefix: str, table_idx: int) -> List[Block]:  # type: ignore[no-untyped-def]
    blocks: List[Block] = []
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            cell_prefix = (
                f"{prefix}/table[{table_idx}]/row[{row_idx}]/cell[{cell_idx}]"
            )
            blocks.extend(_read_paragraphs(cell.paragraphs, part, cell_prefix))
            for nested_idx, nested in enumerate(cell.tables):
                blocks.extend(
                    _read_table(nested, part, cell_prefix, nested_idx)
                )
    return blocks


def read_document(data: bytes) -> DocumentSnapshot:
    """Baca DOCX dari bytes menjadi DocumentSnapshot."""
    doc = Document(BytesIO(data))
    snapshot_id = str(uuid.uuid4())
    parts: dict = {}

    body_blocks: List[Block] = []
    body_blocks.extend(_read_paragraphs(doc.paragraphs, "document", "body"))
    for t_idx, table in enumerate(doc.tables):
        body_blocks.extend(_read_table(table, "document", "body", t_idx))
    parts["document"] = body_blocks

    for s_idx, section in enumerate(doc.sections):
        parts[f"header[{s_idx}]"] = _read_paragraphs(
            section.header.paragraphs, f"header[{s_idx}]", f"header[{s_idx}]"
        )
        parts[f"footer[{s_idx}]"] = _read_paragraphs(
            section.footer.paragraphs, f"footer[{s_idx}]", f"footer[{s_idx}]"
        )

    return DocumentSnapshot(snapshot_id=snapshot_id, parts=parts)


__all__ = ["read_document"]