import io

from docx import Document

from src.docxlint.services.dictionary_service import DictionaryService
from src.docxlint.services.docx_reader import read_document
from src.docxlint.services.language_checker import LanguageChecker


def _synthetic_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("salahh ejaan disini")
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "salahh juga"
    doc.sections[0].header.paragraphs[0].text = "salahh header"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_findings_from_body_table_header():
    svc = DictionaryService.from_data(
        words=("ejaan", "juga", "disini", "salah", "header", "sini", "di")
    )
    snapshot = read_document(_synthetic_docx())
    findings = LanguageChecker(svc).check(snapshot)

    origins = {f.location.part for f in findings}
    assert "document" in origins  # body dan tabel
    assert any(p.startswith("header[") for p in origins)

    salahh_findings = [f for f in findings if f.original_value == "salahh"]
    assert len(salahh_findings) >= 3  # body, tabel, header


def test_location_metadata_shape():
    svc = DictionaryService.from_data(words=("ejaan",))
    snapshot = read_document(_synthetic_docx())
    findings = LanguageChecker(svc).check(snapshot)
    for f in findings:
        assert f.location.snapshot_id
        assert f.location.block_path
        assert f.location.block_type
        assert isinstance(f.location.run_index, int)
        assert f.location.char_start >= 0
        assert f.location.char_end >= f.location.char_start
        assert f.finding_id
        assert f.rule_id
        assert f.category
        assert f.severity
        assert f.application_status == "not_applied"
        assert f.review_status == "pending"