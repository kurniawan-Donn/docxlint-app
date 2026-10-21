import io

from docx import Document

from src.docxlint.services.dictionary_service import DictionaryService
from src.docxlint.services.docx_reader import read_document
from src.docxlint.services.foreign_terms_checker import (
    ForeignTermsChecker,
)
from src.docxlint.services.glossary_service import GlossaryService


def _doc():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("machine learning")
    r.italic = False
    doc.add_paragraph("biasa saja")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_findings_from_real_docx():
    g = GlossaryService.from_data([
        {"term": "machine learning", "italic": True, "category": "istilah_asing"},
    ])
    d = DictionaryService.from_data(words=("biasa", "saja"))
    checker = ForeignTermsChecker(g, d)
    findings = checker.check(read_document(_doc()))
    assert any(f.rule_id == "FOREIGN_001" for f in findings)
    assert any(f.rule_id == "ITALIC_001" for f in findings)


def test_locations_are_non_empty():
    g = GlossaryService.from_data([
        {"term": "machine learning", "italic": True},
    ])
    d = DictionaryService.from_data(words=("biasa", "saja"))
    findings = ForeignTermsChecker(g, d).check(read_document(_doc()))
    for f in findings:
        assert f.location.snapshot_id
        assert f.location.block_path
        assert f.location.block_type
        assert f.location.run_index >= 0