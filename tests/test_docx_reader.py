import io

from docx import Document
from docx.oxml import OxmlElement

from src.docxlint.services.docx_reader import read_document


def test_read_paragraph(minimal_docx_bytes):
    snapshot = read_document(minimal_docx_bytes)
    blocks = snapshot.parts["document"]
    assert any(b.text == "Halo dunia" for b in blocks)


def test_read_table(docx_with_table_bytes):
    snapshot = read_document(docx_with_table_bytes)
    texts = [b.text for b in snapshot.parts["document"]]
    for expected in ("A1", "B1", "A2", "B2"):
        assert expected in texts


def test_read_multiple_runs(docx_with_multiple_runs_bytes):
    snapshot = read_document(docx_with_multiple_runs_bytes)
    para = next(
        b for b in snapshot.parts["document"] if b.text == "Halo dunia"
    )
    assert len(para.runs) == 2
    assert para.runs[0].bold is True
    assert para.runs[1].italic is True
    assert para.runs[0].char_start == 0
    assert para.runs[0].char_end == 5
    assert para.runs[1].char_start == 5
    assert para.runs[1].char_end == 10


def test_simple_run_not_complex(minimal_docx_bytes):
    snapshot = read_document(minimal_docx_bytes)
    para = next(b for b in snapshot.parts["document"] if b.runs)
    assert all(not r.is_complex_run for r in para.runs)


def test_complex_run_detection_break():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("baris satu")
    br = OxmlElement("w:br")
    r._element.append(br)
    buf = io.BytesIO()
    doc.save(buf)
    snapshot = read_document(buf.getvalue())
    para = next(b for b in snapshot.parts["document"] if b.runs)
    assert any(run.is_complex_run for run in para.runs)


def test_header_footer_parts_present(minimal_docx_bytes):
    snapshot = read_document(minimal_docx_bytes)
    assert any(k.startswith("header[") for k in snapshot.parts)
    assert any(k.startswith("footer[") for k in snapshot.parts)