import io

import pytest
from docx import Document


@pytest.fixture
def minimal_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Halo dunia")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def docx_with_table_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Paragraf pertama")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def docx_with_multiple_runs_bytes() -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Halo ")
    r1.bold = True
    r2 = p.add_run("dunia")
    r2.italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()